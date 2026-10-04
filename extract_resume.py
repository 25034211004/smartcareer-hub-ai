
import sys
from pathlib import Path


def extract_resume(file_path):
    path = Path(file_path)

    if not path.is_file():
        raise ValueError("Resume file not found.")

    extension = path.suffix.lower()

    if extension == ".pdf":
        import fitz

        with fitz.open(str(path)) as pdf:
            text = "\n".join(
                page.get_text()
                for page in pdf
            )

        return text.strip()

    if extension == ".docx":
        from docx import Document

        document = Document(str(path))

        paragraphs = [
            p.text
            for p in document.paragraphs
        ]

        # Include text from tables.
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.append(cell.text)

        return "\n".join(paragraphs).strip()

    raise ValueError(
        "Only PDF and DOCX files are supported."
    )


if __name__ == "__main__":
    try:
        if len(sys.argv) != 2:
            raise ValueError(
                "Resume file path is required."
            )

        extracted_text = extract_resume(
            sys.argv[1]
        )

        if not extracted_text:
            raise ValueError(
                "No extractable text found in resume."
            )

        sys.stdout.buffer.write(
            extracted_text.encode("utf-8")
        )

    except Exception as error:
        print(
            str(error),
            file=sys.stderr
        )

        sys.exit(1)
